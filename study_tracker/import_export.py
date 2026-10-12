"""Safe import of SQLite backups and StudyTracker CSV/XLSX exports."""

from __future__ import annotations

import csv
import json
import os
import sqlite3
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

from openpyxl import load_workbook

from .models import COLUMNS

JST = ZoneInfo("Asia/Tokyo")
DATABASE_TABLES = {"sessions", "segments"}


@dataclass
class ImportSummary:
    imported: int = 0
    skipped: int = 0
    aliases_added: int = 0
    warnings: list[str] = field(default_factory=list)


def _database_tables(connection: sqlite3.Connection) -> set[str]:
    return {
        str(row[0])
        for row in connection.execute(
            "SELECT name FROM sqlite_master WHERE type='table'"
        ).fetchall()
    }


def validate_database_file(path: str | Path) -> Path:
    """Verify SQLite integrity and the minimum StudyTracker schema."""
    source = Path(path).expanduser().resolve()
    if not source.is_file():
        raise FileNotFoundError(f"SQLiteファイルが見つかりません: {source}")
    connection = None
    try:
        connection = sqlite3.connect(f"file:{source.as_posix()}?mode=ro", uri=True, timeout=30)
        integrity = connection.execute("PRAGMA integrity_check").fetchone()
        if integrity is None or integrity[0] != "ok":
            raise sqlite3.DatabaseError(f"整合性チェックに失敗しました: {integrity!r}")
        missing = DATABASE_TABLES - _database_tables(connection)
        if missing:
            raise ValueError(
                "StudyTrackerのバックアップではありません。必須テーブルがありません: "
                + ", ".join(sorted(missing))
            )
    except sqlite3.Error as exc:
        raise ValueError(f"SQLiteファイルを読み込めません: {exc}") from exc
    finally:
        if connection is not None:
            connection.close()
    return source


def _unique_backup_path(directory: Path, prefix: str) -> Path:
    stamp = datetime.now().astimezone().strftime("%Y%m%d_%H%M%S_%f")
    directory.mkdir(parents=True, exist_ok=True)
    candidate = directory / f"{prefix}_{stamp}.sqlite3"
    suffix = 1
    while candidate.exists():
        candidate = directory / f"{prefix}_{stamp}_{suffix}.sqlite3"
        suffix += 1
    return candidate


def preserve_database(
    database_path: str | Path,
    backup_directory: str | Path,
    *,
    prefix: str = "study_tracker_before_import",
) -> Path:
    """Create a durable, non-pruned SQLite snapshot of the current database."""
    source_path = validate_database_file(database_path)
    destination = _unique_backup_path(Path(backup_directory).expanduser().resolve(), prefix)
    source = target = None
    try:
        source = sqlite3.connect(str(source_path), timeout=30)
        target = sqlite3.connect(str(destination), timeout=30)
        source.backup(target)
        integrity = target.execute("PRAGMA integrity_check").fetchone()
        if integrity is None or integrity[0] != "ok":
            raise sqlite3.DatabaseError(f"整合性チェックに失敗しました: {integrity!r}")
        target.close()
        target = None
        source.close()
        source = None
        return destination
    except Exception:
        destination.unlink(missing_ok=True)
        raise
    finally:
        if target is not None:
            target.close()
        if source is not None:
            source.close()


def restore_database_snapshot(
    snapshot_path: str | Path,
    destination_path: str | Path,
    backup_directory: str | Path,
) -> Path:
    """Preserve current data, stage and validate the snapshot, then replace atomically.

    The caller must close the live Database connection before this function is called.
    Returns the path of the separately preserved previous database.
    """
    snapshot = validate_database_file(snapshot_path)
    destination = Path(destination_path).expanduser().resolve()
    if snapshot == destination:
        raise ValueError("読み込むバックアップと現在のDBが同じファイルです。")
    if not destination.is_file():
        raise FileNotFoundError(f"現在のStudyTracker DBが見つかりません: {destination}")

    preserved = preserve_database(destination, backup_directory)
    destination.parent.mkdir(parents=True, exist_ok=True)
    stage = destination.with_name(
        f".{destination.name}.{datetime.now().strftime('%Y%m%d%H%M%S%f')}.importing"
    )
    source = target = None
    try:
        source = sqlite3.connect(str(snapshot), timeout=30)
        target = sqlite3.connect(str(stage), timeout=30)
        source.backup(target)
        integrity = target.execute("PRAGMA integrity_check").fetchone()
        if integrity is None or integrity[0] != "ok":
            raise sqlite3.DatabaseError(f"読み込みデータの整合性チェックに失敗しました: {integrity!r}")
        missing = DATABASE_TABLES - _database_tables(target)
        if missing:
            raise ValueError("必須テーブルがありません: " + ", ".join(sorted(missing)))
        target.close()
        target = None
        source.close()
        source = None
        os.replace(stage, destination)
        return preserved
    finally:
        if target is not None:
            target.close()
        if source is not None:
            source.close()
        stage.unlink(missing_ok=True)


def _normalize_cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat(sep=" ")
    return str(value).strip()


def _read_csv(path: Path) -> tuple[list[list[str]], list[tuple[str, str]]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        fieldnames = [str(value or "").strip() for value in (reader.fieldnames or [])]
        if not set(COLUMNS).issubset(fieldnames):
            missing = [column for column in COLUMNS if column not in fieldnames]
            raise ValueError("StudyTracker CSVの必須列がありません: " + ", ".join(missing))
        rows = [
            [_normalize_cell(record.get(column)) for column in COLUMNS]
            for record in reader
        ]
    return rows, []


def _read_xlsx(path: Path) -> tuple[list[list[str]], list[tuple[str, str]]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook["学習記録"] if "学習記録" in workbook.sheetnames else workbook.active
        iterator = sheet.iter_rows(values_only=True)
        header = [_normalize_cell(value) for value in next(iterator, ())]
        indexes = {name: index for index, name in enumerate(header)}
        missing = [column for column in COLUMNS if column not in indexes]
        if missing:
            raise ValueError("StudyTracker XLSXの必須列がありません: " + ", ".join(missing))
        rows: list[list[str]] = []
        for raw in iterator:
            rows.append([
                _normalize_cell(raw[indexes[column]] if indexes[column] < len(raw) else "")
                for column in COLUMNS
            ])

        aliases: list[tuple[str, str]] = []
        if "学習単位エイリアス" in workbook.sheetnames:
            alias_rows = workbook["学習単位エイリアス"].iter_rows(values_only=True)
            alias_header = [_normalize_cell(value) for value in next(alias_rows, ())]
            try:
                title_index = alias_header.index("タイトル")
                doc_index = alias_header.index("Docstring")
            except ValueError:
                title_index = doc_index = -1
            if title_index >= 0:
                for row in alias_rows:
                    title = _normalize_cell(row[title_index] if title_index < len(row) else "")
                    doc = _normalize_cell(row[doc_index] if doc_index >= 0 and doc_index < len(row) else "")
                    if title and doc:
                        aliases.append((title, doc))
        return rows, aliases
    finally:
        workbook.close()


def _resolve_branch_cells(rows: list[list[str]]) -> list[list[str]]:
    resolved: list[list[str]] = []
    for row_index, raw in enumerate(rows, 1):
        row = list(raw)
        for column, value in enumerate(row[:10]):
            if value.strip().lower() == "branch":
                if not resolved:
                    raise ValueError(f"{row_index}行目: 先行行のない branch が見つかりました。")
                row[column] = resolved[-1][column]
        resolved.append(row)
    return resolved


def _parse_date(value: str) -> datetime.date:
    for fmt in ("%Y.%m.%d", "%Y-%m-%d", "%Y/%m/%d"):
        try:
            return datetime.strptime(value.strip(), fmt).date()
        except ValueError:
            continue
    raise ValueError(f"年月日の形式が不正です: {value}")


def _parse_time(value: str) -> datetime.time:
    token = value.strip()
    if not token or token == "--:--:--":
        raise ValueError("開始時刻または終了時刻が未設定です。")
    for fmt in ("%H:%M:%S.%f", "%H:%M:%S", "%H:%M"):
        try:
            return datetime.strptime(token, fmt).time()
        except ValueError:
            continue
    raise ValueError(f"時刻の形式が不正です: {value}")


def _row_payload(values: list[str]) -> dict[str, Any]:
    return {
        "date": values[0],
        "start": values[1],
        "end": values[2],
        "timezone": values[3] or "Asia/Tokyo",
        "category": values[4],
        "subject": values[5],
        "material": values[6],
        "position": values[7],
        "unit": values[8],
        "scope": values[9],
        "actual_scope": "",
        "notes": [line for line in values[10].splitlines() if line.strip()],
        "branch_fields": [],
    }


def import_export_file(database, source_path: str | Path) -> ImportSummary:
    """Append records from a StudyTracker CSV/XLSX export without replacing existing data."""
    path = Path(source_path).expanduser().resolve()
    if not path.is_file():
        raise FileNotFoundError(f"読み込むファイルが見つかりません: {path}")
    if path.suffix.lower() == ".csv":
        rows, aliases = _read_csv(path)
    elif path.suffix.lower() == ".xlsx":
        rows, aliases = _read_xlsx(path)
    else:
        raise ValueError("エクスポートの読み込みは .csv または .xlsx を指定してください。")

    summary = ImportSummary()
    if not rows:
        summary.warnings.append("ファイルに学習記録がありません。")
        return summary

    rows = _resolve_branch_cells(rows)
    prepared: list[tuple[datetime, datetime, dict[str, Any]]] = []
    for index, values in enumerate(rows, 1):
        if not any(value.strip() for value in values):
            summary.skipped += 1
            continue
        try:
            date_value = _parse_date(values[0])
            start_time = _parse_time(values[1])
            end_time = _parse_time(values[2])
            start_at = datetime.combine(date_value, start_time, tzinfo=JST)
            end_at = datetime.combine(date_value, end_time, tzinfo=JST)
            if end_at <= start_at:
                raise ValueError("終了時刻は開始時刻より後である必要があります。")
            prepared.append((start_at, end_at, _row_payload(values)))
        except ValueError as exc:
            summary.skipped += 1
            summary.warnings.append(f"{index}行目を読み飛ばしました: {exc}")

    connection = database.con
    try:
        connection.execute("BEGIN")
        # Each distinct time interval becomes a session; multiple hierarchy rows
        # that share that interval remain grouped together.
        session_ids: dict[tuple[str, str], int] = {}
        for start_at, end_at, payload in prepared:
            key = (start_at.isoformat(), end_at.isoformat())
            if key not in session_ids:
                cursor = connection.execute(
                    "INSERT INTO sessions(started_at,ended_at,active) VALUES(?,?,0)",
                    key,
                )
                session_ids[key] = int(cursor.lastrowid)
            connection.execute(
                "INSERT INTO segments(session_id,start_at,end_at,active,row_json) VALUES(?,?,?,?,?)",
                (
                    session_ids[key],
                    start_at.isoformat(),
                    end_at.isoformat(),
                    1,
                    json.dumps(payload, ensure_ascii=False),
                ),
            )
            summary.imported += 1

        for title, docstring in aliases:
            cursor = connection.execute(
                "INSERT OR IGNORE INTO aliases(title,docstring) VALUES(?,?)",
                (title, docstring),
            )
            if cursor.rowcount > 0:
                summary.aliases_added += 1
        connection.commit()
    except Exception:
        connection.rollback()
        raise

    if summary.skipped:
        summary.warnings.insert(
            0,
            f"{summary.skipped}件は開始・終了日時を正確に復元できないため読み込みませんでした。",
        )
    return summary
